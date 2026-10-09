"""XLSX 引擎：基于 openpyxl 解析上传的工作簿。

连接参数来自 datasource_conns：conn_meta 含 file_path（暂存路径）。
sheet 名作为「表」；列类型由 sync 流程按多行样例推断，也支持用户指定覆盖。
"""

from __future__ import annotations

import os
import time

from openpyxl import load_workbook

from app.core.exceptions import ApiError, ErrorCode
from app.modules.datasource.sync.engines.base import BaseEngine, EngineColumn, EngineTable


class XlsxEngine(BaseEngine):
    def __init__(self, meta: dict, credentials: dict | None = None):
        self.file_path = meta.get("file_path", "")
        if not self.file_path or not os.path.exists(self.file_path):
            raise ApiError(ErrorCode.NOT_FOUND, f"XLSX 文件不存在或已被清理: {self.file_path}", status_code=404)

    def _load(self):
        try:
            return load_workbook(self.file_path, read_only=True, data_only=True)
        except Exception as exc:
            raise ApiError(ErrorCode.VALIDATION, f"无法解析 XLSX 文件: {exc}", status_code=400) from exc

    def test_conn(self) -> dict:
        started = time.time()
        wb = self._load()
        try:
            return {
                "ok": True,
                "latency_ms": int((time.time() - started) * 1000),
                "message": f"XLSX 文件可用，含 {len(wb.sheetnames)} 个工作表",
            }
        finally:
            wb.close()

    def list_tables(self) -> list[EngineTable]:
        wb = self._load()
        try:
            out: list[EngineTable] = []
            for name in wb.sheetnames:
                ws = wb[name]
                rows, _cols = ws.max_row or 0, ws.max_column or 0
                out.append(EngineTable(database_name=os.path.basename(self.file_path), name=name, comment=f"{rows} 行"))
            if not out:
                raise ApiError(ErrorCode.NOT_FOUND, "XLSX 工作簿为空", status_code=400)
            return out
        finally:
            wb.close()

    def columns(self, table: str) -> list[EngineColumn]:
        """返回该 sheet 的列；source_type 为推断出的 Doris 类型（xlsx 无原生类型）。"""
        wb = self._load()
        try:
            if table not in wb.sheetnames:
                raise ApiError(ErrorCode.NOT_FOUND, f"工作簿无此工作表: {table}", status_code=404)
            ws = wb[table]
            rows = list(ws.iter_rows(max_row=1 + 100, values_only=True))
            header, data = rows[0], rows[1:]
            inferred = _infer_from_sample(header, data)
            return [EngineColumn(name=c.name, source_type=c.doris_type) for c in inferred]
        finally:
            wb.close()

    def read_rows(self, table: str, batch_size: int = 1000):
        wb = self._load()
        try:
            if table not in wb.sheetnames:
                raise ApiError(ErrorCode.NOT_FOUND, f"工作簿无此工作表: {table}", status_code=404)
            ws = wb[table]
            batch: list[list] = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i == 0:
                    continue  # 表头
                batch.append(list(row))
                if len(batch) >= batch_size:
                    yield batch
                    batch = []
            if batch:
                yield batch
        finally:
            wb.close()


def _infer_from_sample(header: list, sample_rows: list):
    """根据表头 + 样例行推断列定义（复用 common 的类型推断）。"""
    from app.modules.datasource.sync.common import infer_xlsx_columns

    return infer_xlsx_columns(list(header or []), [list(r) for r in sample_rows if r])